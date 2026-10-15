from pathlib import Path
import json, csv, sys, collections, hashlib, shutil

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
ASSETS = OUT / 'assets'
ASSETS.mkdir(exist_ok=True)
def load(p): return json.loads((ROOT / p).read_text(encoding='utf-8'))
A, B = [load(f'output/resnet18/{m}_seed42.json') for m in ('post','early')]
VA, VB = [load(f'output/resnet18/{m}_val.json') for m in ('post','early')]
AUD = load('dataset/pilot_tier3_4events/audit.json')
READY = load('output/automation_xbd/real_data_readiness.json')
LABELS = ['no_damage','minor','major','destroyed']
CN = ['无损毁','轻度损毁','重度损毁','完全摧毁']
def table(headers, rows):
    return '\n'.join(['| '+' | '.join(map(str,headers))+' |','| '+' | '.join(['---']*len(headers))+' |']+['| '+' | '.join(map(str,r))+' |' for r in rows])
def f(x): return f'{x:.4f}'
def pct(x): return f'{x*100:.2f}%'
def pp(x): return f'{x*100:+.3f}'
def add(title, body): sections.append((title,body.strip()))

if '--figures' in sys.argv:
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import numpy as np
    plt.rcParams.update({'font.family':'Microsoft YaHei','axes.spines.top':False,'axes.spines.right':False,'font.size':10,'figure.dpi':160})
    fig,ax=plt.subplots(1,2,figsize=(10,3.6))
    for d,c,n in [(A,'#245A81','B1 仅灾后'),(B,'#BB6133','B2 灾前与灾后')]:
        h=d['history'];ax[0].plot([x['epoch'] for x in h],[x['train_loss'] for x in h],color=c,label=n)
        ax[1].plot([x['epoch'] for x in h],[x['val']['macro_f1'] for x in h],color=c,label=n)
        ax[1].scatter(d['best_epoch'],d['best_val']['macro_f1'],color=c,s=40,zorder=3)
    ax[0].set(title='训练加权交叉熵',xlabel='训练轮次',ylabel='Loss');ax[1].set(title='验证集四类 Macro F1',xlabel='训练轮次',ylabel='Macro F1',ylim=(0,0.7))
    for a in ax:a.legend();a.grid(alpha=.15);a.set_xticks([1,5,10,15,20])
    fig.tight_layout();fig.savefig(ASSETS/'training_curves.png');plt.close(fig)
    fig,axes=plt.subplots(1,2,figsize=(10,4.1))
    for ax,d,name in zip(axes,[VA,VB],['B1 仅灾后','B2 灾前与灾后']):
        cm=np.array(d['metrics']['confusion']);norm=cm/cm.sum(axis=1,keepdims=True)
        ax.imshow(norm,cmap='Blues',vmin=0,vmax=1)
        for i in range(4):
            for j in range(4):ax.text(j,i,f'{cm[i,j]:,}\n{norm[i,j]*100:.1f}%',ha='center',va='center',color='white' if norm[i,j]>.55 else '#111111',fontsize=9)
        ax.set_xticks(range(4),CN);ax.set_yticks(range(4),CN);ax.set(xlabel='预测类别',ylabel='真实类别',title=name)
    fig.tight_layout();fig.savefig(ASSETS/'confusion_matrices.png');plt.close(fig)
    fig,ax=plt.subplots(figsize=(9,3.2));x=np.arange(4)
    ax.bar(x-.18,[VA['metrics']['per_class'][k]['f1'] for k in LABELS],.36,label='B1',color='#245A81');ax.bar(x+.18,[VB['metrics']['per_class'][k]['f1'] for k in LABELS],.36,label='B2',color='#BB6133')
    ax.set_xticks(x,CN);ax.set(ylabel='F1',ylim=(0,1.05),title='最佳模型在同一验证集上的逐类 F1');ax.legend();ax.grid(axis='y',alpha=.15);fig.tight_layout();fig.savefig(ASSETS/'class_f1.png');plt.close(fig)
    print('Figures ready');sys.exit(0)

# Validate every row of the paired predictions before using their summaries.
pred=[]
for mode in ['post','early']:
    with (ROOT/f'output/resnet18/{mode}_val.csv').open(encoding='utf-8',newline='') as h:pred.append(list(csv.DictReader(h)))
assert len(pred[0])==len(pred[1])==22372
paired=collections.Counter()
for x,y in zip(*pred):
    assert all(x[k]==y[k] for k in ['event','tile_id','building_index','actual'])
    paired[(x['actual']==x['predicted'],y['actual']==y['predicted'])]+=1
assert VA['metrics']==A['best_val'] and VB['metrics']==B['best_val']
assert A['manifest_sha256']==B['manifest_sha256']==hashlib.sha256((ROOT/'dataset/pilot_tier3_4events/manifest_classification.csv').read_bytes()).hexdigest()
sections=[]
add('摘要与实验范围', '''本实验研究在已知建筑物位置的条件下，加入灾前卫星影像是否改善建筑物损毁四分类，以及这一变化需要付出多少计算成本。实验使用官方 xBD／xView2 GeoTIFF 数据中的四个事件子集，在真实建筑标注生成的相同裁剪上，对比 ResNet-18 仅灾后模型 B1 与灾前灾后六通道早期融合模型 B2。两组从配对的随机初始化出发，各训练 20 轮，固定 seed 为 42，以验证集四类 Macro F1 选择最佳权重。

在 Moore 龙卷风验证集的 22,372 栋建筑上，B1 与 B2 的 Macro F1 分别为 0.6012 和 0.6023。B2 提高 0.001042，即约 0.104 个百分点；Balanced Accuracy 提高约 4.493 个百分点，但 Accuracy 下降约 1.716 个百分点。完全摧毁类 F1 上升，轻度和重度损毁类 F1 下降。这说明加入灾前信息改变了类别间的识别取舍，当前主指标收益很小。

两组模型均已完成正式训练、最佳权重恢复及独立验证复核。这里的“独立”指另起评估流程恢复权重，不代表增加了独立数据集；验证集仍用于过模型选择。最终 Nepal 洪水测试集尚未评估。当前结果只支持四事件子集上的单 seed 验证结论，不支持统计显著提升、全量 xBD 性能或跨灾害类型稳定泛化的结论。

本报告日期为 2026 年 10 月 15 日。报告涵盖数据、模型、训练、验证、资源测量和可复现记录；Stage A 建筑定位、预测轮廓端到端系统和多 seed 实验不属于此次已完成内容。''')

add('1 研究问题与实验设计', '''灾后建筑评估需要区分原有建筑外观与灾后变化。仅灾后模型从灾后纹理、屋顶形态及周围环境判断损毁；配对模型额外获得同一位置的灾前图像，理论上可以学习变化线索，但也可能受到配准误差、拍摄条件和背景差异的影响。

RQ1 关注灾前信息是否提高四类 Macro F1。RQ2 关注性能变化与参数量、显存、前向延迟和训练时间之间的关系。类别层面的分析用于判断总体分数是否掩盖不同损毁等级之间的取舍。

核心控制变量包括相同的训练和验证建筑、事件划分、裁剪窗口、图像缩放、像素处理、ResNet-18 主干系列、优化器、训练轮次、batch、随机种子和 checkpoint 选择规则。主要输入差异是 B1 使用灾后 3 通道，B2 使用灾前加灾后 6 通道；为适配输入，B2 第一层卷积相应扩展。模型参数量因此并非严格相等，但增加部分仅位于首层。

建筑裁剪位置来自真实建筑标注。该设计隔离分类问题，避免自动建筑定位误差干扰主比较。它不能直接代表从整幅影像自动发现建筑并评估损毁的端到端系统性能。

执行顺序为数据校验与审计、真实数据小规模试跑、正式单 seed 成对训练、恢复最佳权重评估验证集、汇总指标和成本。小规模试跑使用 128 个训练样本、64 个验证样本、batch 8、1 轮，只验证输入输出和资源可用性。本文所有主结果来自使用完整选定训练集和验证集的正式 20 轮实验。''')

add('2 数据来源与子集组成', '''数据来自官方 xBD／xView2 GeoTIFF 影像和 JSON 建筑标注 [1]。原始下载包含 part-aa 至 part-af 六个分卷，它们共同组成一个压缩归档，不是六个独立数据集。六个正式分卷到齐后按顺序流式读取，整体 SHA1 校验通过，再只提取 tier3 中四个选定事件。没有把归档内全部灾害事件用于训练，也没有额外生成完整合并归档或全量解压。

归档校验值为 6eae3baddf86796c15638682a6432e3e6223cb39。四事件提取写入 5,144 个文件，共 16,359,343,887 字节，约 16.36 GB 或 15.24 GiB，包括 2,572 幅影像与 2,572 份标签，对应 1,286 对灾前灾后 tile。建筑样本是标注中的单栋建筑，与 tile 或影像文件的计数不同。

'''+table(['划分','灾害事件','建筑样本数'],[
['训练','joplin-tornado',14696],['训练','lower-puna-volcano',2856],['验证','moore-tornado',22372],['测试','nepal-flooding',41582],['合计','四事件',81506]])+'''

训练总数 17,552，验证总数 22,372，测试总数 41,582。按事件划分，而不是在同一灾害事件内部随机拆分建筑，有助于减少相邻场景和灾害背景泄漏。这是本项目自定义的四事件划分，不能直接当作官方 benchmark 的标准划分进行横向排名。文件夹名称中的 pilot 表示子集沿用的目录名称，不表示正式训练仍被限制为 128 个样本。

类别依次为 no_damage、minor、major、destroyed，对应标签 ID 0、1、2、3。未分类标签不被强行映射为某一种损毁程度；本次共过滤 3,479 条未知损毁标签，保留 81,506 条有效四分类记录。''')

add('3 类别分布与数据质量',table(['类别','训练数','训练占比','验证数','验证占比','测试数'],[[CN[i],READY['per_split_class_counts']['train'][k],pct(READY['per_split_class_counts']['train'][k]/17552),READY['per_split_class_counts']['val'][k],pct(READY['per_split_class_counts']['val'][k]/22372),READY['per_split_class_counts']['test'][k]] for i,k in enumerate(LABELS)])+'''

验证集无损毁建筑占约 86.95%，明显高于训练集约 59.83%。验证集轻度、重度和完全摧毁类分别仅 886、449 和 1,584 栋，说明不能只用总体 Accuracy 判断模型。测试集数量仅用于数据组成说明，没有生成测试预测或测试性能指标。

结构审计显示缺失组件为 0，重复建筑组合标识为 0，事件、tile 和影像跨划分重叠为 0。manifest 实际引用 2,220 幅影像，均可读取，尺寸均为 1024×1024，三通道，存储类型为 int16，nodata 为 -99，坐标参考系统为 EPSG:4326。建筑框越界计数为 0，影像审计错误为 0。

2,220 是有效建筑清单实际用到的影像数，2,572 是提取出的总影像文件数，两者统计对象不同。前期人工查看了灾前灾后与建筑 polygon 叠加图，在总览尺度确认建筑位置对应；该检查不等于对全部样本做精确配准误差测量，也不能排除局部偏移、阴影、云层或标注歧义。

训练、验证和测试的数据完整性检查可以在模型训练前完成。测试标签参与清单审计和类别数量统计，并不意味着测试集已用于模型选优；本次权重训练仅使用 train，checkpoint 选择仅使用 val。''')

add('4 建筑样本构造与预处理', '''每条 manifest 记录关联同一建筑的灾前影像路径、灾后影像路径、事件、tile、建筑编号、损毁标签及灾后建筑边界框。加载器按需从 GeoTIFF 读取建筑窗口，不预先保存所有裁剪图，减少重复磁盘存储。

具体窗口构造先计算边界框宽 w 与高 h，再设正方形窗口边长为 ceil(max(w,h)+2×16)，以建筑框中心为窗口中心，并对起点向下取整。灾前和灾后使用完全相同的像素窗口。这个规则以标注中对应位置已经可配对为前提，不包含额外的图像配准网络。

读取使用双线性重采样，输出 128×128。窗口允许越过影像边界；nodata 或负数像素转换为 0，再截断到 [0,255] 并除以 255。最终每张 RGB 建筑图为 float32 的 3×128×128 张量。int16 是原始存储格式，不意味着按 16 位最大值归一化。

B1 输入为灾后张量。B2 按灾前在前、灾后在后的顺序在通道维拼接，得到 6×128×128。训练没有额外使用 ImageNet 均值方差标准化、旋转、翻转或颜色增强，故不能将这些方法写成已完成设置。

当前加载器对两组均读取灾前和灾后裁剪，再由输入函数决定是否使用灾前张量。因此 B1 的实际数据读取流程也包含灾前影像读取；训练耗时比较不是优化后的纯灾后系统与配对系统的完整端到端 I/O 成本比较。后文的前向 benchmark 则只测模型输入张量常驻设备后的网络计算。

样本由真实建筑框提供，裁剪中可能保留邻近建筑、道路或植被。这有助于提供上下文，也可能使模型依赖背景线索；当前没有 polygon 掩膜抠图，也没有进行背景消融。''')

add('5 模型结构与初始化', '''正式主干采用 Torchvision ResNet-18 [2]，不使用预训练权重，分类头从标准任务输出改为四类。B1 与 B2 共用同一结构系列，B2 仅扩展首层输入通道。SmallCNN 是早期软件检查模型，不产生本报告的正式结果。

'''+table(['阶段','具体运算','输出形状 不含 batch'],[
['输入','B1 灾后 RGB 或 B2 灾前灾后拼接','3×128×128 或 6×128×128'],['首层','7×7 卷积 64通道 步长2 padding3','64×64×64'],['激活与池化','BatchNorm ReLU 3×3最大池化 步长2','64×32×32'],['残差阶段1','2个 BasicBlock 每块2个3×3卷积','64×32×32'],['残差阶段2','2个 BasicBlock 首块下采样','128×16×16'],['残差阶段3','2个 BasicBlock 首块下采样','256×8×8'],['残差阶段4','2个 BasicBlock 首块下采样','512×4×4'],['全局池化','Adaptive Average Pooling','512×1×1'],['分类头','Flatten 和 Linear 512→4','4个 logits']])+'''

基本残差块为两层 3×3 卷积，每层配 BatchNorm，第一层后使用 ReLU，卷积支路与 shortcut 相加后再使用 ReLU。通道或空间尺寸变化时，shortcut 使用 1×1 卷积与 BatchNorm 投影。残差形式可写为 y=ReLU(F(x)+S(x))。输出为四个未归一化 logits，训练直接交给交叉熵；评估以 argmax 选类别，计算置信度时才使用 softmax。

为控制初始化差异，两组先以相同种子创建完整三通道网络。设 B1 首层权重为 W，B2 首层设为 concat(W,W)/2，其余层保持相同。因此输入的灾前灾后张量完全相同时，两组首层输出一致。这个初始化规则不代表训练后的两组权重仍相同。

B1 有 11,178,564 个参数，B2 有 11,187,972 个参数。增加的 9,408 个参数等于 64×3×7×7，仅约 0.0842%。六通道输入没有使整个主干参数翻倍。''')

add('6 训练配置与评价指标',table(['配置项','正式实验设置'],[
['训练数据','17,552 栋 全部使用 max_train=0'],['验证数据','22,372 栋 全部使用 max_val=0'],['随机种子','42 仅一组 seed'],['轮次与 batch','每模型20轮 batch=32 workers=0'],['优化器','AdamW lr=0.001 weight_decay=0.01'],['其他优化器设置','PyTorch默认 betas=(0.9,0.999) eps=1e-8'],['学习率调度','CosineAnnealingLR T_max=20 eta_min=0'],['损失函数','按训练集类别频次计算的加权交叉熵'],['最佳权重','验证集四类Macro F1严格提升保存 平分保留较早轮次'],['计算精度','float32 未使用自动混合精度'],['确定性','固定Python NumPy Torch种子 使用确定性算法']])+'''

类别权重按逆训练频次计算，再除以四类权重的均值。四类权重按标签顺序分别为 0.214451、1.004980、2.184443、0.596125。权重仅使用训练标签计算，不能由验证或测试标签重新估计。该加权目标减少多数类对训练损失的主导，但可能改变 precision 与 recall 的平衡。

对类别 c，Precision=TP/(TP+FP)，Recall=TP/(TP+FN)，F1=2PR/(P+R)。Macro F1 是固定四类 F1 的算术平均；Accuracy 是全部预测正确数除以样本数；Balanced Accuracy 是有真实样本的类别召回率均值。本次验证集四类均存在，因此它是四类 Recall 的平均。所有指标按建筑汇总，并非先计算每个 tile 后再平均。

训练每轮评估验证集，训练结束后重新加载最佳 checkpoint。另起评估流程检查 manifest 哈希、类别映射和权重来源，输出逐建筑 CSV、指标 JSON 和混淆矩阵。两组恢复后的指标均与训练记录 best_val 完全一致。复核验证了实现一致性，但不会消除使用验证集选优带来的选择偏差。''')

add('7 总体结果与逐类性能',table(['验证指标','B1 仅灾后','B2 灾前灾后','B2减B1 百分点'],[[k,f(VA['metrics'][key]),f(VB['metrics'][key]),pp(VB['metrics'][key]-VA['metrics'][key])] for k,key in [('Accuracy','accuracy'),('Macro F1','macro_f1'),('Balanced Accuracy','balanced_accuracy')]])+'''

B1 最佳权重为第 5 轮，B2 为第 9 轮；两组都完成了 20 轮预算。最佳轮次不同符合相同验证选优规则，并不是两组训练预算不同。以下均为这两个最佳权重在同一验证集上的结果。

'''+table(['类别','模型','Precision','Recall','F1','样本数'],[[CN[i],name,f(d['metrics']['per_class'][k]['precision']),f(d['metrics']['per_class'][k]['recall']),f(d['metrics']['per_class'][k]['f1']),d['metrics']['per_class'][k]['support']] for i,k in enumerate(LABELS) for name,d in [('B1',VA),('B2',VB)]])+'''

![逐类 F1](assets/class_f1.png)

B2 的完全摧毁类 F1 由 0.6893 升至 0.7394，提高约 5.012 个百分点；但无损毁、轻度、重度的 F1 分别下降约 1.008、1.113 和 2.474 个百分点。四类变化平均后，整体 Macro F1 仅轻微上升。仅报告主指标差值会掩盖这种类别间抵消。''')

add('8 混淆矩阵与错误取舍', '''混淆矩阵中行表示真实类别、列表示预测类别。图中每格同时标注建筑数量和在该真实类别内的比例，行归一化颜色可减少多数类计数过大对视觉比较的干扰。

![混淆矩阵](assets/confusion_matrices.png)

'''+table(['真实类别','模型','预测无损毁','预测轻度','预测重度','预测摧毁'],[[CN[i],name,*d['metrics']['confusion'][i]] for i in range(4) for name,d in [('B1',VA),('B2',VB)]])+'''

B2 将更多无损毁建筑预测为轻度或重度损毁：这两类误报合计由 672 增至 1,477；同时将无损毁误判为完全摧毁的数量由 847 降至 463。轻度类 Recall 由 34.09% 升至 46.95%，重度类 Recall 由 50.11% 升至 67.04%，但相关预测的 Precision 下降，故两类 F1 并未随 Recall 同步提高。

完全摧毁类 Precision 由 55.51% 升至 67.84%，Recall 则由 90.91% 降至 81.25%。B2 的“摧毁”预测更精确，但漏掉更多真实摧毁建筑；其中真实摧毁被预测为重度的数量由 44 增至 218。因此不能将较高的摧毁类 F1 解释为救灾场景下所有错误成本都更低。

''' + f'逐建筑 CSV 已按事件、tile、建筑编号和真实标签核对，22,372 条记录完全对应。两模型都正确 {paired[(True,True)]:,} 条，仅 B1 正确 {paired[(True,False)]:,} 条，仅 B2 正确 {paired[(False,True)]:,} 条，两者都错误 {paired[(False,False)]:,} 条。B2 相对 B1 净减少 384 个正确预测，与 Accuracy 下降一致。此处为配对描述统计，没有进行显著性检验或置信区间估计。')

add('9 训练过程与收敛表现', '''![训练曲线](assets/training_curves.png)

两组训练加权交叉熵整体下降。B1 从第 1 轮约 0.9094 降至第 20 轮约 0.0095，B2 从约 0.9355 降至约 0.0113。训练损失低并不代表跨事件验证表现持续改善。

B1 在第 5 轮达到最佳 Macro F1，之后未超越；B2 在第 9 轮达到最佳，第 12 轮接近该水平。最后一轮的 Macro F1 分别为 0.5109 和 0.4978，均低于各自最佳值。若错误地使用最后一轮权重比较，会得到不同且较差的结果。

验证曲线波动明显，例如 B2 第 6 轮降至约 0.1196，随后又恢复。这提示训练动态或跨事件分布差异值得进一步研究。训练损失下降而验证表现不持续提升，与泛化不足或过拟合相容，但现有结果不能定位为单一原因，也不能在未检查时归因于程序错误、图像配准或 BatchNorm。

后续诊断可以在训练和验证范围内检查各轮混淆矩阵、学习率、类别预测分布及代表性图像。当前不依据测试结果改变训练轮次、模型或超参数。完整逐轮记录见附录。''')

costrows=[['参数量',f"{A['params']:,}",f"{B['params']:,}",'+0.084%'],['参数存储 MiB',f(A['parameter_bytes']/2**20),f(B['parameter_bytes']/2**20),'+0.084%'],['训练及逐轮验证总时长 小时',f(A['training_seconds']/3600),f(B['training_seconds']/3600),f"{(B['training_seconds']/A['training_seconds']-1)*100:+.2f}%"],['训练与验证峰值 allocated MiB',f(A['training_peak_allocated_bytes']/2**20),f(B['training_peak_allocated_bytes']/2**20),f"{(B['training_peak_allocated_bytes']/A['training_peak_allocated_bytes']-1)*100:+.2f}%"]]
add('10 资源开销与测量边界',table(['项目','B1','B2','相对变化'],costrows)+'''

'''+table(['模型','batch','批次延迟 ms','每样本摊销 ms','吞吐 样本每秒','峰值 allocated MiB'],[[name,x['batch_size'],f(x['batch_latency_ms']),f(x['amortized_ms_per_image']),f'{x["images_per_second"]:.1f}',f(x['peak_allocated_bytes']/2**20)] for name,d in [('B1',A),('B2',B)] for x in d['inference_benchmark']])+'''

硬件为 NVIDIA GeForce RTX 4060 Laptop GPU。软件记录为 Python 3.12.15、PyTorch 2.7.1+cu128、Torchvision 0.22.1+cu128、NumPy 2.0.2、Rasterio 1.4.3。两组累计训练及逐轮验证约 8.70 小时，未把数据准备、独立评估和报告生成计入。

前向 benchmark 在恢复最佳权重后运行，使用 GPU 常驻 float32 合成输入张量，预热 20 次，测量 100 次并做 CUDA 同步，分别测 batch 1 和 32。它只测网络前向，不含真实 GeoTIFF 读取、裁剪、CPU 到 GPU 传输、建筑定位、结果导出或地图渲染。合成张量在这里仅用于计时，不是分类性能数据。

B2 的 batch 1 和 batch 32 批次延迟分别增加约 1.16% 和 0.84%。这些差异来自一次测量，可能受到设备功耗、温度及测量波动影响，不能宣称稳定速度差异。B2 的记录训练总时长略低约 0.49%，同样不构成“B2 更快”的可靠结论。

显存数值为 PyTorch 分配器报告的 peak allocated memory，不是任务管理器中的整卡显存，也不包含全部驱动、缓存或其他进程占用。MiB 按 2^20 字节换算。参数存储是参数张量字节数，不等于含元数据和缓冲区的 checkpoint 文件大小。总体上本次首层扩展成本较小，但完整部署还需考虑获取、配对和读取灾前影像的额外成本。''')

add('11 讨论与结论边界', '''对 RQ1，当前证据显示灾前输入带来轻微的总体 Macro F1 上升，同时改变了类别的 precision 与 recall。由于差值仅约 0.0010、只有一个随机种子、且分数来自用于选优的单一验证事件，不能宣称灾前信息显著或稳定优于仅灾后输入，也不能据此断言灾前信息没有价值。

对 RQ2，六通道早期融合只扩展首层，参数与测得前向成本增幅有限。但非常小的主指标收益仍不足以证明额外数据准备在实际部署中值得；不同任务对漏报和误报的代价不同，应结合类别表现和业务目标评价。当前没有完成代价敏感评估。

本实验的主要限制包括：仅四事件子集；仅 seed42；验证只有 Moore 龙卷风，不能分离事件特征与灾害类型影响；未评估封存的 Nepal 洪水测试；训练从头初始化且没有增强，结论不覆盖预训练和其他融合架构；真实建筑框降低了定位难度，不能代表自动端到端系统；类别不平衡和分布差异影响总体指标；精确配准及逐图错例成因尚未系统测量；资源 benchmark 不覆盖完整数据处理成本。

建筑之间还存在同一 tile 内的空间关联，22,372 栋验证建筑不能简单视为完全独立的 22,372 次灾害实验。未来若计算不确定性，应考虑按 tile 或事件分组，而不是仅按建筑独立抽样。当前没有报告 p 值、置信区间或多 seed 标准差。

现阶段最稳妥的结论是：在本次固定四事件划分、随机初始化 ResNet-18 和单 seed 条件下，早期融合与仅灾后模型的验证 Macro F1 基本接近，早期融合改善了完全摧毁类 F1 和平均类别召回率，同时牺牲了总体准确率及部分损毁类 F1。该结果为后续复现实验和错误分析提供了依据。''')

add('12 已完成工作与后续实验',table(['工作项','截至报告日期的状态'],[
['六分卷下载和整体校验','完成'],['四事件提取 清单与图像审计','完成'],['真实数据小规模两模式试跑','完成 仅检查软件和资源'],['正式ResNet-18 B1与B2各20轮','完成 seed42'],['恢复最佳权重验证与配对记录核查','完成'],['验证指标 混淆矩阵和资源统计','完成'],['最终封存测试集性能评估','尚未运行'],['三个固定seed重复实验','尚未运行'],['逐图错例成因与大小质量分组分析','尚未完成'],['Stage A定位和端到端damage map','尚未实现本轮实验'],['本地新增实现同步GitHub','报告生成前仍存在未提交修改']])+'''

建议先核查有代表性的轻度重度混淆、无损毁误报及摧毁漏报样本，使用验证数据检查图像质量、配准与上下文。若资源允许，在相同协议下增加 seed43 和 seed44，并报告三 seed 的均值与样本标准差。任何改变都需要单独记录为新实验，不能覆盖当前结果。

模型结构、预处理和评估方案冻结后，再对 B1 与 B2 分别进行一次封存测试评估，并将测试结果与验证选优分数分开报告。如果测试结果不理想，也不应回头利用该测试集反复调参。Stage A 和其他融合结构属于后续扩展，不是对当前四事件单 seed 主比较完成状态的必要追认。

本报告是已完成实验的详细记录，最终课程项目仍需要把测试结果、必要的复现或错例分析、贡献说明和引用整理进最终提交。数据和大模型文件按仓库规则保留在本地，代码、协议和可分享的汇总材料可经检查后同步到仓库。''')

epochrows=[]
for x,y in zip(A['history'],B['history']):epochrows.append([x['epoch'],f(x['train_loss']),f(x['val']['macro_f1']),f(y['train_loss']),f(y['val']['macro_f1'])])
add('附录 A 完整训练轮次记录',table(['轮次','B1训练损失','B1验证Macro F1','B2训练损失','B2验证Macro F1'],epochrows)+'''

本表为日志数据的四位小数展示，图表和复核使用 JSON 中的完整精度。训练损失是按类别权重归一化累计的加权交叉熵，不应与未加权损失或准确率直接比较。最佳权重是 B1 第5轮与 B2 第9轮，最终轮次只用于呈现训练轨迹。''')

sources=[('训练配置 历史和成本','output/resnet18/post_seed42.json；output/resnet18/early_seed42.json'),('最佳模型验证结果','output/resnet18/post_val.json；output/resnet18/early_val.json'),('逐建筑配对预测','output/resnet18/post_val.csv；output/resnet18/early_val.csv'),('验证汇总','output/resnet18/frozen_val_comparison.json'),('归档与标签审计','dataset/pilot_tier3_4events/extraction_summary.json；audit.json'),('影像和分组审计','dataset/pilot_tier3_4events/geotiff_audit.json；output/automation_xbd/real_data_readiness.json'),('模型与预处理实现','stage_b_damage/train_baseline.py；stage_b_damage/dataset.py'),('实验协议','reports/PHASE3_PROTOCOL.md')]
add('附录 B 复现证据与参考文献', '''主要证据文件以项目目录 D:\\DDA4220\\DDA4220 为根。报告不重新训练模型，也不运行最终测试。生成报告时核对了两组训练数据哈希、恢复权重指标及全部验证预测的配对关系。

'''+ '\n\n'.join(f'{i+1}. {a}：{b}' for i,(a,b) in enumerate(sources))+'''

归档 SHA1：6eae3baddf86796c15638682a6432e3e6223cb39

manifest SHA256：39b802dd01776e2f136dba3c12ba7ffb54a52eac173ca306500a45f74651a2a0

B1 checkpoint SHA256：2eb308f6d4d9ecc20cb54f82f3d6845d8f451f20a6b01a4ae28d88e9a6d0f551

B2 checkpoint SHA256：6f88a11cc007028324d223af5bb0ffbcfd57e79a49e0eb58aea8749981a6cab3

参考文献

[1] Gupta R, Hosfelt R, Sajeev S, et al. xBD A Dataset for Assessing Building Damage from Satellite Imagery. 2019. https://arxiv.org/abs/1911.09296

[2] He K, Zhang X, Ren S, Sun J. Deep Residual Learning for Image Recognition. 2015. https://arxiv.org/abs/1512.03385

论文用于说明数据与网络的来源，本文具体实现、划分和性能以本地实验记录为准；未与文献原始 benchmark 直接比较。''')

TITLE='灾前灾后卫星影像建筑损毁分类实验报告'
SUB='xBD 四事件子集上的 ResNet18 单种子受控对比'
md='# '+TITLE+'\n\n'+SUB+'\n\n报告日期：2026年10月15日\n\n'+'\n\n'.join('## '+t+'\n\n'+b for t,b in sections)+'\n'
(OUT/'实验报告.md').write_text(md,encoding='utf-8')
(OUT/'report_sections.json').write_text(json.dumps(sections,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'verification.json').write_text(json.dumps({'validation_rows':22372,'prediction_pairing':'all matched','best_metrics_restored':'exact match both models','manifest_hash':'matched','paired_correctness':{str(k):v for k,v in paired.items()},'test_evaluated':False},indent=2),encoding='utf-8')

if '--markdown-only' in sys.argv:print('Markdown and verification ready');sys.exit(0)
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
import re
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.65);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0);s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
normal=doc.styles['Normal'];normal.font.size=Pt(10);normal.paragraph_format.line_spacing=1.22;normal.paragraph_format.space_after=Pt(7)
doc.styles['Heading 1'].font.size=Pt(17);doc.styles['Heading 1'].paragraph_format.space_after=Pt(12)
doc.styles['Title'].font.size=Pt(23)
doc.core_properties.title=TITLE;doc.core_properties.subject=SUB;doc.core_properties.author='DDA4220 项目组'
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=foot.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
def paragraph(s):
 p=doc.add_paragraph(s);return p
def make_table(lines):
 rows=[[s.strip() for s in line.strip('|').split('|')] for line in lines if not re.match(r'^\|[\s|:-]+\|$',line)]
 t=doc.add_table(rows=0,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=[17.4/len(rows[0])]*len(rows[0])
 if len(rows[0])==2:widths=[4.8,12.6]
 if len(rows[0])==3:widths=[3,10.5,3.9]
 if rows[0][0]=='阶段':widths=[3,8.9,5.5]
 if len(rows[0])==6:widths=[3,1.5,3.2,3.2,3.2,3.3]
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for j,(c,txt) in enumerate(zip(cells,row)):
   c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=c.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.05
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(row)<=3 else WD_ALIGN_PARAGRAPH.CENTER
   r=p.add_run(txt);r.font.size=Pt(8.5)
   pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'70');el.set(qn('w:type'),'dxa');margins.append(el)
   pr.append(margins)
   if i==0 or i%2==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DDE7EF' if i==0 else 'F5F7F9');pr.append(sh)
   if i==0:r.bold=True
  trpr=t.rows[-1]._tr.get_or_add_trPr();avoid=OxmlElement('w:cantSplit');trpr.append(avoid)
  if i==0:
   rep=OxmlElement('w:tblHeader');trpr.append(rep)
 paragraph('')
for idx,(title,body) in enumerate(sections):
 if idx:doc.add_page_break()
 else:
  doc.add_paragraph(TITLE,'Title');doc.add_paragraph(SUB,'Subtitle');paragraph('报告日期 2026年10月15日    验证结果版    seed42')
 doc.add_heading(title,level=1)
 blocks=body.split('\n\n')
 for block in blocks:
  if block.startswith('|'):make_table(block.splitlines())
  elif block.startswith('!['):
   label,path=re.match(r'!\[(.*?)\]\((.*?)\)',block).groups();p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.add_run().add_picture(str(OUT/path),width=Cm(17.1));p=paragraph(label);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(8)
  else:paragraph(block)
doc.save(OUT/'实验报告.docx')
print('Created',OUT/'实验报告.docx','sections',len(sections),'characters',len(md))
